import numpy as np
import pytest
from pypdf import PdfWriter

from backend.pipeline import chunker


class TestCosineSimilarity:
    def test_identical_vectors_returns_one(self):
        vec = np.array([1.0, 2.0, 3.0])
        assert chunker.cosine_similarity(vec, vec) == pytest.approx(1.0)

    def test_orthogonal_vectors_returns_zero(self):
        a = np.array([1.0, 0.0])
        b = np.array([0.0, 1.0])
        assert chunker.cosine_similarity(a, b) == pytest.approx(0.0)

    def test_zero_vector_returns_zero_without_error(self):
        a = np.array([0.0, 0.0])
        b = np.array([1.0, 1.0])
        assert chunker.cosine_similarity(a, b) == 0.0


class TestSplitIntoSentences:
    def test_turkish_ordinal_period_does_not_split_sentence(self):
        text = (
            "50.000 Avro'ya kadar üçüncü kişilere yapılan para transferlerinde "
            "1. Grup C Sınıfı ile 2. Grup C Sınıfı yetkilileri müştereken yetkilidir."
        )
        sentences = chunker.split_into_sentences(text)
        assert sentences == [text]

    def test_genuinely_separate_sentences_still_split(self):
        text = "Sistem çöktü. Ekip müdahale etti."
        sentences = chunker.split_into_sentences(text)
        assert sentences == ["Sistem çöktü.", "Ekip müdahale etti."]

    def test_madde_numarasi_does_not_split_from_following_sentence(self):
        text = "Madde 2.9 - Avro-TL Dönüşümü. Avrodan TL'ye çevrim için TCMB kuru esas alınır."
        sentences = chunker.split_into_sentences(text)
        assert sentences == [
            "Madde 2.9 - Avro-TL Dönüşümü.",
            "Avrodan TL'ye çevrim için TCMB kuru esas alınır.",
        ]


class TestLocateSentenceOffsets:
    def test_finds_sequential_sentences(self):
        raw_text = "Bir. Iki. Uc."
        sentences = ["Bir.", "Iki.", "Uc."]
        offsets = chunker.locate_sentence_offsets(raw_text, sentences)
        assert offsets == [(0, 4), (5, 9), (10, 13)]

    def test_finds_sentences_across_newlines(self):
        raw_text = "Birinci satir.\nIkinci satir.\nUcuncu satir."
        sentences = ["Birinci satir.", "Ikinci satir.", "Ucuncu satir."]
        offsets = chunker.locate_sentence_offsets(raw_text, sentences)
        assert offsets == [(0, 14), (15, 28), (29, 42)]

    def test_missing_sentence_returns_none_pair(self):
        raw_text = "Sadece bu var."
        sentences = ["Sadece bu var.", "Bu metinde yok."]
        offsets = chunker.locate_sentence_offsets(raw_text, sentences)
        assert offsets == [(0, 14), (None, None)]

    def test_empty_sentences_returns_empty_list(self):
        assert chunker.locate_sentence_offsets("herhangi bir metin", []) == []


class TestCreateChunks:
    def test_splits_when_similarity_below_threshold(self):
        sentences = ["A.", "B.", "C."]
        similarities = [0.9, 0.1]  # A-B benzer, B-C degil
        chunks = chunker.create_chunks(sentences, similarities, threshold=0.36)
        assert chunks == ["A. B.", "C."]

    def test_keeps_all_together_when_similarity_above_threshold(self):
        sentences = ["A.", "B.", "C."]
        similarities = [0.9, 0.8]
        chunks = chunker.create_chunks(sentences, similarities, threshold=0.36)
        assert chunks == ["A. B. C."]

    def test_single_sentence_returns_single_chunk(self):
        chunks = chunker.create_chunks(["Tek cumle."], [], threshold=0.36)
        assert chunks == ["Tek cumle."]

    def test_empty_sentences_returns_empty_list(self):
        assert chunker.create_chunks([], [], threshold=0.36) == []


class TestCreateChunksWithRanges:
    def test_splits_when_similarity_below_threshold(self):
        sentences = ["A.", "B.", "C."]
        similarities = [0.9, 0.1]
        ranges = chunker.create_chunks_with_ranges(sentences, similarities, threshold=0.36)
        assert ranges == [("A. B.", 0, 1), ("C.", 2, 2)]

    def test_keeps_all_together_when_similarity_above_threshold(self):
        sentences = ["A.", "B.", "C."]
        similarities = [0.9, 0.8]
        ranges = chunker.create_chunks_with_ranges(sentences, similarities, threshold=0.36)
        assert ranges == [("A. B. C.", 0, 2)]

    def test_single_sentence_returns_single_range(self):
        ranges = chunker.create_chunks_with_ranges(["Tek cumle."], [], threshold=0.36)
        assert ranges == [("Tek cumle.", 0, 0)]

    def test_empty_sentences_returns_empty_list(self):
        assert chunker.create_chunks_with_ranges([], [], threshold=0.36) == []

    def test_create_chunks_stays_a_thin_wrapper(self):
        """create_chunks'in mevcut kontratini (dus metin listesi) korudugunu dogrular."""
        sentences = ["A.", "B.", "C."]
        similarities = [0.9, 0.1]
        assert chunker.create_chunks(sentences, similarities, threshold=0.36) == ["A. B.", "C."]


class TestLoadDocument:
    def test_loads_txt_as_plain_text(self, tmp_path):
        f = tmp_path / "doc.txt"
        f.write_text("Merhaba dunya.", encoding="utf-8")
        assert chunker.load_document(str(f)) == "Merhaba dunya."

    def test_loads_md_as_plain_text(self, tmp_path):
        f = tmp_path / "doc.md"
        f.write_text("# Baslik\n\nIcerik.", encoding="utf-8")
        assert chunker.load_document(str(f)) == "# Baslik\n\nIcerik."

    def test_loads_pdf_without_error(self, tmp_path):
        f = tmp_path / "doc.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(f, "wb") as out:
            writer.write(out)

        result = chunker.load_document(str(f))
        assert isinstance(result, str)


class TestLoadDocumentDocx:
    def test_extracts_paragraph_text_from_docx(self, tmp_path):
        import docx
        document = docx.Document()
        document.add_paragraph("Birinci paragraf.")
        document.add_paragraph("Ikinci paragraf.")
        path = tmp_path / "test.docx"
        document.save(str(path))

        text = chunker.load_document(str(path))
        assert "Birinci paragraf." in text
        assert "Ikinci paragraf." in text


class TestLoadDocumentHtml:
    def test_strips_tags_scripts_and_styles(self, tmp_path):
        html = (
            "<html><head><style>body{color:red}</style></head>"
            "<body><script>alert(1)</script><p>Merhaba dunya.</p></body></html>"
        )
        path = tmp_path / "test.html"
        path.write_text(html, encoding="utf-8")

        text = chunker.load_document(str(path))
        assert "Merhaba dunya." in text
        assert "alert" not in text
        assert "color:red" not in text


class TestFindSourceFiles:
    def test_filters_supported_extensions_only(self, tmp_path):
        (tmp_path / "a.txt").write_text("x", encoding="utf-8")
        (tmp_path / "b.md").write_text("x", encoding="utf-8")
        (tmp_path / "c.pdf").write_bytes(b"%PDF-1.4")
        (tmp_path / "ignore.json").write_text("{}", encoding="utf-8")

        found = chunker.find_source_files(str(tmp_path))

        assert [f.split(chunker.os.sep)[-1] for f in found] == ["a.txt", "b.md", "c.pdf"]

    def test_empty_dir_returns_empty_list(self, tmp_path):
        assert chunker.find_source_files(str(tmp_path)) == []


class TestGetEmbeddings:
    def test_calls_model_encode(self):
        class FakeModel:
            def encode(self, sentences):
                return [[float(len(s))] for s in sentences]

        result = chunker.get_embeddings(["a", "bb", "ccc"], FakeModel())
        assert result == [[1.0], [2.0], [3.0]]


class TestChunkDocument:
    def test_integrates_embeddings_with_similarity_split(self, tmp_path, monkeypatch):
        """Regresyon testi: get_embeddings tanımsızken chunk_document NameError ile patlıyordu
        (get_embeddings hiçbir testte cagrilmadigi icin fark edilmemisti)."""
        f = tmp_path / "doc.txt"
        f.write_text("Bir. Iki. Uc.", encoding="utf-8")

        monkeypatch.setattr(chunker, "split_into_sentences", lambda text: ["Bir.", "Iki.", "Uc."])
        monkeypatch.setattr(
            chunker,
            "get_embeddings",
            lambda sentences, model: np.array([[1.0, 0.0], [1.0, 0.0], [0.0, 1.0]]),
        )

        chunks = chunker.chunk_document(str(f), model=None, threshold=0.36)

        assert chunks == [("Bir. Iki.", "satır 1"), ("Uc.", "satır 1")]


class TestSlugify:
    def test_strips_extension_and_keeps_alnum_underscore(self):
        assert chunker.slugify("mimari_v1.txt") == "mimari_v1"

    def test_replaces_unsafe_characters_with_underscore(self):
        assert chunker.slugify("my doc (v2).pdf") == "my_doc_v2_"


class TestChunkIdFor:
    def test_produces_stable_per_file_id(self):
        assert chunker.chunk_id_for("mimari_v1.txt", 0) == "mimari_v1_ch0001"
        assert chunker.chunk_id_for("mimari_v1.txt", 9) == "mimari_v1_ch0010"

    def test_different_sources_never_collide(self):
        assert chunker.chunk_id_for("wikipedia_deadlock.txt", 0) == "wikipedia_deadlock_ch0001"


class TestAlreadyIngestedSources:
    def test_returns_empty_set_when_file_missing(self, tmp_path):
        missing = tmp_path / "chunks.jsonl"
        assert chunker.already_ingested_sources(str(missing)) == set()

    def test_returns_distinct_source_ids_from_existing_file(self, tmp_path):
        f = tmp_path / "chunks.jsonl"
        f.write_text(
            '{"chunk_id": "a_ch0001", "source_id": "a.txt", "location": "chunk_index_0", "text": "x", "quote": "x..."}\n'
            '{"chunk_id": "a_ch0002", "source_id": "a.txt", "location": "chunk_index_1", "text": "y", "quote": "y..."}\n'
            '{"chunk_id": "b_ch0001", "source_id": "b.txt", "location": "chunk_index_0", "text": "z", "quote": "z..."}\n',
            encoding="utf-8",
        )
        assert chunker.already_ingested_sources(str(f)) == {"a.txt", "b.txt"}


class TestSaveToJsonl:
    def test_writes_expected_schema(self, tmp_path):
        output_file = tmp_path / "out" / "chunks.jsonl"
        records = [("Bu bir test cumlesidir ve devami vardir.", "doc.txt", "satır 1")]

        chunker.save_to_jsonl(records, str(output_file))

        lines = output_file.read_text(encoding="utf-8").strip().splitlines()
        assert len(lines) == 1

        import json
        data = json.loads(lines[0])
        assert data["chunk_id"] == "doc_ch0001"
        assert data["source_id"] == "doc.txt"
        assert data["location"] == "satır 1"
        assert data["quote"] == "Bu bir test cumlesidir ve..."
        assert "ingested_at" in data

    def test_ingested_at_is_iso_format_utc(self, tmp_path):
        from datetime import datetime
        output_file = tmp_path / "chunks.jsonl"
        chunker.save_to_jsonl([("metin", "doc.txt", "satır 1")], str(output_file))

        import json
        data = json.loads(output_file.read_text(encoding="utf-8").strip())
        # datetime.fromisoformat basariyla parse edebiliyorsa format gecerlidir
        parsed = datetime.fromisoformat(data["ingested_at"])
        assert parsed.tzinfo is not None


class TestSaveToJsonlIncremental:
    def test_appends_to_existing_file_with_per_source_chunk_ids(self, tmp_path):
        output = tmp_path / "chunks.jsonl"
        output.write_text(
            '{"chunk_id": "a_ch0001", "source_id": "a.txt", "location": "chunk_index_0", "text": "old", "quote": "old..."}\n',
            encoding="utf-8",
        )

        chunker.save_to_jsonl([("yeni parca", "b.txt", "satır 1")], str(output))

        lines = output.read_text(encoding="utf-8").strip().split("\n")
        assert len(lines) == 2
        import json
        first = json.loads(lines[0])
        second = json.loads(lines[1])
        assert first["chunk_id"] == "a_ch0001"  # eski satir dokunulmadan kaldi
        assert second["chunk_id"] == "b_ch0001"  # yeni dosya kendi 1'den numaralandirmasiyla basladi
        assert second["source_id"] == "b.txt"

    def test_multiple_chunks_same_source_number_sequentially(self, tmp_path):
        output = tmp_path / "chunks.jsonl"
        chunker.save_to_jsonl([("birinci", "c.txt", "satır 1"), ("ikinci", "c.txt", "satır 2")], str(output))

        import json
        lines = [json.loads(line) for line in output.read_text(encoding="utf-8").strip().split("\n")]
        assert lines[0]["chunk_id"] == "c_ch0001"
        assert lines[1]["chunk_id"] == "c_ch0002"


class FakePdfPage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


class FakePdfReader:
    def __init__(self, pages):
        self.pages = pages


class TestLoadPdfWithPageOffsets:
    def test_computes_page_start_offsets(self, monkeypatch, tmp_path):
        fake_pages = [FakePdfPage("AB"), FakePdfPage("CDE")]
        monkeypatch.setattr(chunker, "PdfReader", lambda path: FakePdfReader(fake_pages))
        f = tmp_path / "doc.pdf"
        f.write_bytes(b"%PDF-1.4")

        text, page_offsets = chunker.load_pdf_with_page_offsets(str(f))

        assert text == "AB\nCDE"
        assert page_offsets == [0, 3]

    def test_handles_empty_page_text(self, monkeypatch, tmp_path):
        fake_pages = [FakePdfPage(None), FakePdfPage("X")]
        monkeypatch.setattr(chunker, "PdfReader", lambda path: FakePdfReader(fake_pages))
        f = tmp_path / "doc.pdf"
        f.write_bytes(b"%PDF-1.4")

        text, page_offsets = chunker.load_pdf_with_page_offsets(str(f))

        assert text == "\nX"
        assert page_offsets == [0, 1]


class TestFormatLineLocation:
    def test_single_line_range(self):
        raw = "Birinci satir.\nIkinci satir.\nUcuncu satir."
        start = raw.index("Ikinci")
        end = start + len("Ikinci satir.")
        assert chunker.format_line_location(raw, start, end) == "satır 2"

    def test_multi_line_range(self):
        raw = "Bir.\nIki.\nUc.\nDort."
        start = 0
        end = raw.index("Uc.") + len("Uc.")
        assert chunker.format_line_location(raw, start, end) == "satır 1-3"


class TestFormatPageLocation:
    def test_single_page(self):
        page_offsets = [0, 20, 45]
        assert chunker.format_page_location(5, 15, page_offsets) == "sayfa 1"

    def test_multi_page_range(self):
        page_offsets = [0, 20, 45]
        assert chunker.format_page_location(15, 30, page_offsets) == "sayfa 1-2"

    def test_offset_at_exact_page_boundary_counts_as_new_page(self):
        page_offsets = [0, 20, 45]
        assert chunker.format_page_location(20, 20, page_offsets) == "sayfa 2"


class TestResolveLocation:
    def test_missing_offsets_falls_back_to_chunk_index(self):
        result = chunker._resolve_location(None, None, 3, "herhangi bir metin", None)
        assert result == "chunk_index_3"

    def test_uses_line_location_when_no_page_offsets(self):
        raw = "Bir.\nIki."
        result = chunker._resolve_location(5, 9, 0, raw, None)
        assert result == "satır 2"

    def test_uses_page_location_when_page_offsets_given(self):
        result = chunker._resolve_location(25, 30, 0, "herhangi bir metin", [0, 20, 45])
        assert result == "sayfa 2"
